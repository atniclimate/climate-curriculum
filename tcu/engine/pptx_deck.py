"""Write a course deck as .pptx from the same slide list the HTML deck uses.

Starts from tcu/templates/pptx/tcu-deck.pptx (16:9, dark master). The ATNI
fonts (League Spartan, Lexend Deca) fall back to system fonts on machines
where they are not installed.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

import mdlite

TEMPLATE = Path(__file__).resolve().parent.parent / 'templates' / 'pptx' / 'tcu-deck.pptx'
WHITE, TEXT, MUTED, RED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE8, 0xEC, 0xF0), RGBColor(0x9A, 0xA4, 0xB5), RGBColor(0xF2, 0x6B, 0x5E)
FALLBACK = ('Fonts: League Spartan and Lexend Deca (ATNI Climate design system). Where they are not installed, PowerPoint '
            'shows a fallback font; nothing else changes.')
MAX_CHARS = 900


def _text(md):
    s = re.sub(r'\{\{socket:[^}]+\}\}', '[Local lesson slot]', md)
    s = re.sub(r'\{\{view:[^}]+\}\}', '[Data view: see the minisite place page]', s)
    s = re.sub(r'^\|?\s*:?-{2,}.*$', '', s, flags=re.M)
    s = re.sub(r'^\|\s*(.*?)\s*\|\s*$', lambda m: ' | '.join(c.strip() for c in m.group(1).split('|')), s, flags=re.M)
    return re.sub(r'\n{3,}', '\n\n', mdlite.plain(s)).strip()


def _box(slide, x, y, w, h, text, size, color, font, bold=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.name, r.font.bold = Pt(size), color, font, bold
    return tb


def write(slides, dest, course_title, site_name):
    prs = Presentation(str(TEMPLATE))
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    m = int(W * .06)
    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        _box(sl, m, int(H * .06), W - 2 * m, int(H * .07), s['kicker'].upper(), 12, RED, 'League Spartan', True)
        _box(sl, m, int(H * .13), W - 2 * m, int(H * .16), s['title'], 34 if s['hall'] else 28, WHITE, 'League Spartan', True)
        body = _text(s['md'])
        rest = ''
        if len(body) > MAX_CHARS:
            cut = body.rfind('\n', 0, MAX_CHARS)
            cut = cut if cut > 200 else MAX_CHARS
            body, rest = body[:cut] + '\n(continued in the notes)', body[cut:]
        if body:
            _box(sl, m, int(H * .31), W - 2 * m, int(H * .62), body, 20 if s['big'] else 14, TEXT, 'Lexend Deca')
        if i == 0:
            _box(sl, m, int(H * .88), W - 2 * m, int(H * .06), 'DRAFT FOR TCU REVIEW', 12, RED, 'League Spartan', True)
        notes = _text(s['notes'])
        if rest:
            notes = 'Continued from the slide:\n' + rest.strip() + ('\n\n' + notes if notes else '')
        if i == 0:
            notes = (notes + '\n\n' if notes else '') + FALLBACK
        sl.notes_slide.notes_text_frame.text = notes or ' '
    prs.core_properties.title = f'{course_title} ({site_name})'
    prs.core_properties.author = 'Patrick A. Freeland (module author), with ATNI Climate'
    prs.core_properties.comments = 'Template draft for TCU review. ' + FALLBACK
    prs.save(str(dest))
