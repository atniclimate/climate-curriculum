"""Print the text of a document as plain Markdown-ish text. Read-only.

Usage: python -P extract.py <file> [--max-chars N]
Supports .docx .pptx .pdf .md .txt .html .htm. Other types print a one-line notice.
"""
import sys, re, html
from pathlib import Path


def docx_text(p):
    import docx
    d = docx.Document(str(p))
    out = []
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or '').lower()
        if style.startswith('heading'):
            level = ''.join(ch for ch in style if ch.isdigit()) or '2'
            out.append('#' * min(int(level) + 1, 6) + ' ' + t)
        else:
            out.append(t)
    for ti, table in enumerate(d.tables):
        out.append(f'\n[table {ti + 1}]')
        for row in table.rows:
            out.append(' | '.join(c.text.strip().replace('\n', ' ') for c in row.cells))
    return '\n\n'.join(out)


def pptx_text(p):
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f'## Slide {i}')
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    out.append(t)
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            if n:
                out.append('Notes: ' + n)
    return '\n\n'.join(out)


def pdf_text(p):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    return '\n\n'.join(f'[page {i}]\n' + (pg.extract_text() or '') for i, pg in enumerate(r.pages, 1))


def html_text(p):
    s = p.read_text(encoding='utf-8', errors='replace')
    s = re.sub(r'(?is)<(script|style).*?</\1>', '', s)
    s = re.sub(r'(?s)<[^>]+>', ' ', s)
    return re.sub(r'[ \t]+', ' ', html.unescape(s))


def main():
    p = Path(sys.argv[1])
    limit = int(sys.argv[sys.argv.index('--max-chars') + 1]) if '--max-chars' in sys.argv else 0
    ext = p.suffix.lower()
    fn = {'.docx': docx_text, '.pptx': pptx_text, '.pdf': pdf_text, '.html': html_text, '.htm': html_text}.get(ext)
    if fn:
        try:
            text = fn(p)
        except Exception as e:  # corrupt or password-protected files are common in the archive
            print(f'[unreadable {ext}: {p.name}: {type(e).__name__}: {e}]')
            return
    elif ext in ('.md', '.txt'):
        text = p.read_text(encoding='utf-8', errors='replace')
    else:
        print(f'[unsupported type {ext}: {p.name}]')
        return
    text = re.sub(r'\n{3,}', '\n\n', text)
    if limit and len(text) > limit:
        text = text[:limit] + f'\n\n[truncated at {limit} of {len(text)} chars]'
    sys.stdout.reconfigure(encoding='utf-8')
    print(text)


if __name__ == '__main__':
    main()
